from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from a_stock_agents.config import load_settings
from a_stock_agents.recommendation_store import recommendation_store


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "runtime" / "exports" / "每日推荐数据整理.xlsx"


def read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            item = json.loads(line)
        except json.JSONDecodeError:
            continue
        if isinstance(item, dict):
            rows.append(item)
    return rows


def date_of(value: object) -> str:
    text = str(value or "")
    return text[:10] if len(text) >= 10 else ""


def money(value: object) -> float | None:
    try:
        return round(float(value), 4)
    except (TypeError, ValueError):
        return None


def main() -> int:
    settings = load_settings()
    store = recommendation_store(settings.runtime)
    recommendations = store.history()
    results = read_jsonl(settings.runtime / "agents" / "backtest" / "data" / "results.jsonl")
    result_by_id = {str(row.get("recommendation_id")): row for row in results}

    rows = []
    for rec in recommendations:
        result = result_by_id.get(str(rec.get("recommendation_id")))
        basis = rec.get("basis") or {}
        rows.append(
            {
                "推荐日期": date_of(rec.get("created_at") or rec.get("recommendation_date")),
                "生成时间": rec.get("created_at"),
                "股票代码": rec.get("code"),
                "股票名称": rec.get("name"),
                "策略": "平衡多因子" if rec.get("model") == "balanced" else rec.get("model"),
                "模型": rec.get("model"),
                "参考价": money(rec.get("reference_price")),
                "目标卖出日": date_of(rec.get("target_date")),
                "模型评分": money(rec.get("score")),
                "当日涨跌幅%": money(basis.get("change_pct")),
                "成交金额": money(basis.get("amount")),
                "换手率%": money(basis.get("turnover_pct")),
                "市盈率": money(basis.get("pe_ratio")),
                "相关新闻数": basis.get("news_count"),
                "新闻情绪": money(basis.get("news_sentiment")),
                "短期走势强度": money(basis.get("momentum_z")),
                "资金活跃度": money(basis.get("liquidity_z")),
                "回测收益率%": money(result.get("return_pct")) if result else None,
                "回测状态": "盈利" if result and result.get("profitable") is True else ("亏损" if result and result.get("profitable") is False else "待回测"),
                "卖出收盘价": money(result.get("exit_close_price")) if result else None,
                "回测方法": result.get("method") if result else "",
                "推荐ID": rec.get("recommendation_id"),
            }
        )
    rows.sort(key=lambda row: str(row["推荐日期"]), reverse=True)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "每日推荐明细"
    headers = list(rows[0].keys()) if rows else [
        "推荐日期", "生成时间", "股票代码", "股票名称", "策略", "模型", "参考价", "目标卖出日", "模型评分",
        "回测收益率%", "回测状态", "推荐ID"
    ]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        status = row[headers.index("回测状态")].value if "回测状态" in headers else None
        if status == "盈利":
            row[headers.index("回测收益率%")].font = Font(color="00A66A")
        elif status == "亏损":
            row[headers.index("回测收益率%")].font = Font(color="D9364F")
    for col_idx, header in enumerate(headers, 1):
        max_len = max(len(str(header)), *(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(2, ws.max_row + 1)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 36)

    summary = wb.create_sheet("汇总")
    total = len(rows)
    evaluated = [row for row in rows if row["回测状态"] in ("盈利", "亏损")]
    wins = [row for row in evaluated if row["回测状态"] == "盈利"]
    losses = [row for row in evaluated if row["回测状态"] == "亏损"]
    compound = 1.0
    for row in evaluated:
        compound *= 1 + float(row["回测收益率%"] or 0) / 100
    summary_rows = [
        ("指标", "数值"),
        ("推荐记录数", total),
        ("已回测记录数", len(evaluated)),
        ("盈利次数", len(wins)),
        ("亏损次数", len(losses)),
        ("胜率", len(wins) / len(evaluated) if evaluated else None),
        ("累计复利收益率", compound - 1 if evaluated else None),
        ("数据来源", str(settings.runtime / "agents" / "recommendation" / "data" / "recommendations.db")),
    ]
    for item in summary_rows:
        summary.append(item)
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    summary["B6"].number_format = "0.00%"
    summary["B7"].number_format = "0.00%"
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 80

    wb.save(OUTPUT_PATH)
    print(str(OUTPUT_PATH))
    print(f"rows={total}, evaluated={len(evaluated)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
