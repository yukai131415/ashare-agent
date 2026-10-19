from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "runtime" / "research_backtests" / "daily_balanced_close_baostock"
TRADES_CSV = BASE / "trades_99_from_cache.csv"
REPORT_JSON = BASE / "report_99_from_cache.json"
OUTPUT_PATH = BASE / "平衡多因子99天回测明细.xlsx"


HEADERS = [
    ("buy_date", "买入日期"),
    ("sell_date", "卖出日期"),
    ("code", "股票代码"),
    ("name", "股票名称"),
    ("buy_close", "买入收盘价"),
    ("sell_close", "卖出收盘价"),
    ("score", "模型评分"),
    ("return_pct", "单笔收益率%"),
]


def as_float(value: object) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def main() -> int:
    if not TRADES_CSV.exists():
        raise FileNotFoundError(f"缺少 99 天回测明细：{TRADES_CSV}")
    with TRADES_CSV.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    report = json.loads(REPORT_JSON.read_text(encoding="utf-8")) if REPORT_JSON.exists() else {}

    wb = Workbook()
    ws = wb.active
    ws.title = "99天回测明细"
    ws.append([label for _, label in HEADERS])
    for row in rows:
        ws.append([
            row.get("buy_date"),
            row.get("sell_date"),
            row.get("code"),
            row.get("name"),
            as_float(row.get("buy_close")),
            as_float(row.get("sell_close")),
            as_float(row.get("score")),
            as_float(row.get("return_pct")),
        ])

    blue = PatternFill("solid", fgColor="1F4E78")
    white_bold = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = blue
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=5).number_format = "0.00"
        ws.cell(row=row, column=6).number_format = "0.00"
        ws.cell(row=row, column=7).number_format = "0.000000"
        ws.cell(row=row, column=8).number_format = "0.00"
        value = ws.cell(row=row, column=8).value
        ws.cell(row=row, column=8).font = Font(color="00A66A" if value and value > 0 else "D9364F")
    for col_idx in range(1, len(HEADERS) + 1):
        max_len = max(len(str(ws.cell(row=row, column=col_idx).value or "")) for row in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 3, 12), 24)

    summary = wb.create_sheet("汇总")
    summary_rows = [
        ("指标", "数值"),
        ("策略", report.get("strategy", "平衡多因子 balanced")),
        ("买入规则", report.get("buy", "当天收盘买入")),
        ("卖出规则", report.get("sell", "第二个交易日收盘卖出")),
        ("交易笔数", len(rows)),
        ("首个买入日", rows[0].get("buy_date") if rows else ""),
        ("最后买入日", rows[-1].get("buy_date") if rows else ""),
        ("复利总收益率", (as_float(report.get("compound_return_pct")) or 0) / 100),
        ("单笔收益简单相加", (as_float(report.get("simple_sum_return_pct")) or 0) / 100),
        ("胜率", (as_float(report.get("win_rate_pct")) or 0) / 100),
        ("数据源", report.get("data_source", "BaoStock 历史日线缓存")),
        ("口径说明", report.get("factor_note", "")),
    ]
    for item in summary_rows:
        summary.append(item)
    for cell in summary[1]:
        cell.fill = blue
        cell.font = white_bold
    for coord in ("B8", "B9", "B10"):
        summary[coord].number_format = "0.00%"
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 110
    summary["B12"].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(OUTPUT_PATH)
    print(OUTPUT_PATH)
    print(f"rows={len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
